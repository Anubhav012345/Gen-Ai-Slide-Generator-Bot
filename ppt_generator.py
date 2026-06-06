from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


class PPTHandler:

    def set_background(self, slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 0, 0)

    def create_title_slide(self, prs, title, topic, teacher_name):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.set_background(slide)

        tb = slide.shapes.add_textbox(
            Inches(1), Inches(2), Inches(11), Inches(1.2)
        )
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 0)

        tb2 = slide.shapes.add_textbox(
            Inches(1), Inches(3.4), Inches(11), Inches(0.8)
        )
        p2 = tb2.text_frame.paragraphs[0]
        p2.text = f"Topic: {topic}"
        p2.font.size = Pt(22)
        p2.font.color.rgb = RGBColor(255, 255, 255)

        tb3 = slide.shapes.add_textbox(
            Inches(1), Inches(4.4), Inches(11), Inches(0.7)
        )
        p3 = tb3.text_frame.paragraphs[0]
        p3.text = f"By: {teacher_name}"
        p3.font.size = Pt(18)
        p3.font.color.rgb = RGBColor(200, 200, 200)

    def add_content_slide(self, prs, heading, points):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.set_background(slide)

        # Heading box
        hb = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9)
        )
        ht = hb.text_frame
        hp = ht.paragraphs[0]
        hp.text = heading
        hp.font.size = Pt(26)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(255, 255, 0)

        # Content box — tall enough to show full text
        cb = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.7)
        )
        tf = cb.text_frame
        tf.word_wrap = True

        clean_points = [
            pt.strip() for pt in points if len(pt.strip()) >= 5
        ]

        for idx, pt in enumerate(clean_points[:5]):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•  {pt}"
            p.font.size = Pt(17)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.space_after = Pt(14)

    def create_points_presentation(
        self, points_list, title, topic, teacher_name
    ):
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        self.create_title_slide(prs, title, topic, teacher_name)

        for item in points_list:
            heading = item.get("heading", topic)
            points = item.get("points", [])
            if points:
                self.add_content_slide(prs, heading, points)

        output_file = f"{topic.replace(' ', '_')}_points.pptx"
        prs.save(output_file)
        return output_file