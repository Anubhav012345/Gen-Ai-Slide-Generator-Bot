from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


class AIPresentationGenerator:
    def create_presentation(self, extracted_data, title, topic, teacher_name):
        prs = Presentation()
        prs.slide_width = Inches(14)
        prs.slide_height = Inches(7.5)

        # Title slide
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_background(slide, prs, RGBColor(0, 0, 0))

        txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(12), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.color.rgb = RGBColor(255, 255, 0)
        p.font.bold = True

        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.2), Inches(12), Inches(1))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = f"Topic: {topic}"
        p2.font.size = Pt(26)
        p2.font.color.rgb = RGBColor(255, 255, 255)

        txBox3 = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(12), Inches(1))
        tf3 = txBox3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = f"By: {teacher_name}"
        p3.font.size = Pt(20)
        p3.font.color.rgb = RGBColor(200, 200, 200)

        # Content slides
        for point_data in extracted_data:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._set_background(slide, prs, RGBColor(0, 0, 0))

            txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(13), Inches(1.2))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = point_data.get("heading", "")
            p.font.size = Pt(28)
            p.font.color.rgb = RGBColor(255, 255, 0)
            p.font.bold = True

            content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(13), Inches(5.2))
            ctf = content_box.text_frame
            ctf.word_wrap = True

            points = point_data.get("points", [])
            for i, point in enumerate(points):
                if i == 0:
                    cp = ctf.paragraphs[0]
                else:
                    cp = ctf.add_paragraph()
                cp.text = f"•  {point}"
                cp.font.size = Pt(20)
                cp.font.color.rgb = RGBColor(255, 255, 255)
                cp.space_after = Pt(6)

        output_path = f"{topic.replace(' ', '_')}_points.pptx"
        prs.save(output_path)
        return output_path

    def _set_background(self, slide, prs, color):
        background = slide.shapes.add_shape(
            1, 0, 0, prs.slide_width, prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = color
        background.line.fill.background()
        sp = background._element
        slide.shapes._spTree.remove(sp)
        slide.shapes._spTree.insert(2, sp)
